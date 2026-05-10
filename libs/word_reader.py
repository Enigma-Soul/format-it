from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from libs.config import FormatConfig
from libs.fonts import FontResolver
from libs.models import DocumentStructure, FontInfo, InlineRun, ParagraphNode, ParagraphRole


class WordReader:
    def __init__(self, config: FormatConfig, tmp_dir: Path) -> None:
        self._config = config
        self._tmp_dir = tmp_dir

    def read(self, path: Path) -> DocumentStructure:
        doc = Document(str(path))
        image_dir = self._tmp_dir / f"{path.stem}_images"
        image_dir.mkdir(parents=True, exist_ok=True)

        paragraphs: list[ParagraphNode] = []
        img_counter = 0

        # Extract images from relationships
        image_map = self._extract_all_images(doc, image_dir)

        for i, para in enumerate(doc.paragraphs):
            runs = self._extract_runs(para)
            text = para.text.strip()

            # Check for images in this paragraph's XML
            para_images: list[Path] = []
            for drawing in para._element.iter(qn("wp:inline")):
                blip = drawing.find(qn("a:blip"))
                if blip is not None:
                    rId = blip.get(qn("r:embed"))
                    if rId and rId in image_map:
                        para_images.append(image_map[rId])

            # Also check w:drawing (anchored images)
            for drawing in para._element.iter(qn("wp:anchor")):
                blip = drawing.find(qn("a:blip"))
                if blip is not None:
                    rId = blip.get(qn("r:embed"))
                    if rId and rId in image_map:
                        para_images.append(image_map[rId])

            node = ParagraphNode(
                index=i,
                role=ParagraphRole.UNKNOWN,
                text=text,
                runs=runs,
                images=para_images,
                raw_xml_element=para._element,
            )
            paragraphs.append(node)

        # Handle tables
        table_idx = 0
        for table in doc.tables:
            table_md = self._table_to_markdown(table)
            node = ParagraphNode(
                index=len(paragraphs),
                role=ParagraphRole.TABLE,
                text=table_md,
            )
            paragraphs.append(node)
            table_idx += 1

        return DocumentStructure(
            paragraphs=paragraphs,
            images=list(image_map.values()),
            source_file=path,
        )

    def _extract_runs(self, paragraph) -> list[InlineRun]:
        runs: list[InlineRun] = []
        for run in paragraph.runs:
            font = run.font
            size_pt = self._emu_to_pt(font.size) if font.size else None
            if size_pt is None:
                size_pt = 0.0
            font_name = font.name or ""
            color_rgb = None
            if font.color and font.color.rgb:
                color_rgb = str(font.color.rgb)

            font_info = FontInfo(
                name=FontResolver.normalize(font_name),
                size_pt=size_pt,
                bold=font.bold or False,
                italic=font.italic or False,
                color_rgb=color_rgb,
            )
            runs.append(InlineRun(
                text=run.text,
                font=font_info,
            ))
        return runs

    def _extract_all_images(self, doc: Document, image_dir: Path) -> dict[str, Path]:
        image_map: dict[str, Path] = {}
        counter = 0
        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.reltype:
                try:
                    blob = rel.target_part.blob
                    ext = self._guess_image_ext(rel.target_part.content_type)
                    path = image_dir / f"img_{counter}{ext}"
                    path.write_bytes(blob)
                    image_map[rel_id] = path
                    counter += 1
                except Exception:
                    continue
        return image_map

    @staticmethod
    def _table_to_markdown(table) -> str:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if len(rows) >= 1:
            col_count = len(table.rows[0].cells)
            rows.insert(1, "| " + " | ".join(["---"] * col_count) + " |")
        return "\n".join(rows)

    @staticmethod
    def _emu_to_pt(emu) -> float:
        return emu / 12700.0

    @staticmethod
    def _guess_image_ext(content_type: str) -> str:
        mapping = {
            "image/png": ".png",
            "image/jpeg": ".jpg",
            "image/gif": ".gif",
            "image/bmp": ".bmp",
            "image/tiff": ".tiff",
        }
        return mapping.get(content_type, ".png")
