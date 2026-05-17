from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Mm, Pt

from libs.config import FormatConfig
from libs.fonts import FontResolver
from libs.models import DocumentStructure, HEADING_ROLES, ParagraphNode, ParagraphRole


class WordWriter:
    def __init__(self, config: FormatConfig) -> None:
        self._config = config

    def write(self, doc: DocumentStructure, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()

        self._setup_page_layout(document)

        for node in doc.paragraphs:
            if node.role == ParagraphRole.TABLE:
                self._write_table(document, node.text)
                continue
            if node.role == ParagraphRole.PAGE_NUMBER:
                continue
            self._write_paragraph(document, node)

        self._add_page_numbers(document)

        document.save(str(path))
        return path

    def _setup_page_layout(self, document: Document) -> None:
        section = document.sections[0]
        section.page_width = Mm(self._config.paper.width_mm)
        section.page_height = Mm(self._config.paper.height_mm)
        section.top_margin = Mm(self._config.paper.margin_top_mm)
        section.bottom_margin = Mm(self._config.paper.margin_bottom_mm)
        section.left_margin = Mm(self._config.paper.margin_left_mm)
        section.right_margin = Mm(self._config.paper.margin_right_mm)
        if self._config.paper.footer_distance_mm > 0:
            section.footer_distance = Mm(self._config.paper.footer_distance_mm)

    def _write_paragraph(self, document: Document, node: ParagraphNode) -> None:
        if node.images:
            for img_path in node.images:
                try:
                    if img_path.exists():
                        document.add_picture(str(img_path), width=Mm(156))
                except Exception:
                    pass
            return

        if not node.text.strip():
            return

        para = document.add_paragraph()

        # Set alignment
        is_centered = node.role in (ParagraphRole.TITLE, ParagraphRole.SUBTITLE)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_centered else WD_ALIGN_PARAGRAPH.LEFT
        font_spec = self._config.font_for_role(node.role, node.heading_level)

        # Set paragraph format
        pf = para.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(self._config.body_line_spacing_pt)
        pf.space_before = Pt(self._config.space_before_pt)
        pf.space_after = Pt(self._config.space_after_pt)

        if node.role == ParagraphRole.BODY or font_spec.indent:
            pf.first_line_indent = Pt(font_spec.size_pt * 2)
        else:
            pf.first_line_indent = Pt(0)

        # Write text with font handling
        self._write_runs(para, node.text, font_spec, node.role)

    def _write_runs(self, para, text: str, font_spec, role: ParagraphRole) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn as _qn

        resolved_font = FontResolver.resolve(font_spec.name)
        if not resolved_font:
            resolved_font = font_spec.alt_name or "SimSun"

        parts = self._split_text_for_numbers(text)

        for part_text, is_number in parts:
            run = para.add_run(part_text)
            run.font.size = Pt(font_spec.size_pt)
            run.font.bold = font_spec.bold
            number_font = self._config.number_font
            if is_number:
                run.font.name = number_font
            else:
                run.font.name = resolved_font
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(_qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            if is_number:
                rFonts.set(_qn("w:ascii"), number_font)
                rFonts.set(_qn("w:hAnsi"), number_font)
                rFonts.set(_qn("w:eastAsia"), resolved_font)
            else:
                rFonts.set(_qn("w:ascii"), number_font)
                rFonts.set(_qn("w:hAnsi"), number_font)
                rFonts.set(_qn("w:eastAsia"), resolved_font)

    @staticmethod
    def _split_text_for_numbers(text: str) -> list[tuple[str, bool]]:
        parts: list[tuple[str, bool]] = []
        current = ""
        in_number = False
        for ch in text:
            is_digit = ch.isdigit() or ch in ".,-%‰°"
            if is_digit != in_number and current:
                parts.append((current, in_number))
                current = ""
                in_number = is_digit
            if not current:
                in_number = is_digit
            current += ch
        if current:
            parts.append((current, in_number))
        return parts

    def _write_table(self, document: Document, md_text: str) -> None:
        lines = [l for l in md_text.split("\n") if l.strip()]
        if not lines:
            return
        # Parse markdown table
        rows_data: list[list[str]] = []
        for line in lines:
            if line.strip().startswith("|") and "---" not in line:
                cells = [c.strip() for c in line.strip("|").split("|")]
                rows_data.append(cells)

        if not rows_data:
            return

        col_count = len(rows_data[0])
        table = document.add_table(rows=len(rows_data), cols=col_count)
        table.style = "Table Grid"

        for i, row_data in enumerate(rows_data):
            for j, cell_text in enumerate(row_data):
                if j < col_count:
                    table.rows[i].cells[j].text = cell_text

    def _add_page_numbers(self, document: Document) -> None:
        pn = self._config.page_number
        if not pn.enabled:
            return

        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        # Enable odd/even page headers & footers in document settings
        settings = document.settings.element
        even_odd = settings.find(qn("w:evenAndOddHeaders"))
        if even_odd is None:
            even_odd = OxmlElement("w:evenAndOddHeaders")
            settings.append(even_odd)

        section = document.sections[0]

        # Odd page footer (default)
        odd_footer = section.footer
        odd_footer.is_linked_to_previous = False
        self._write_page_number_para(
            odd_footer.paragraphs[0] if odd_footer.paragraphs else odd_footer.add_paragraph(),
            pn.odd_align,
            pn,
            qn,
        )

        # Even page footer
        sect_pr = section._sectPr
        even_footer_ref = sect_pr.find(qn("w:footerReference[@type='even']"))
        if even_footer_ref is None:
            from lxml import etree
            from docx.opc.part import Part
            from docx.opc.packuri import PackURI

            ftr_tag = OxmlElement("w:ftr")
            p_el = OxmlElement("w:p")
            self._build_page_number_xml(p_el, pn.even_align, pn, qn, OxmlElement)
            ftr_tag.append(p_el)
            even_xml_bytes = etree.tostring(ftr_tag, xml_declaration=True, encoding="UTF-8", standalone=True)

            even_partname = PackURI(odd_footer.part.partname.replace("/footer", "/footerEven"))

            even_part = Part(
                even_partname,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml",
                even_xml_bytes,
                odd_footer.part.package,
            )
            rel = document.part.relate_to(
                even_part,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer",
            )

            even_ref = OxmlElement("w:footerReference")
            even_ref.set(qn("w:type"), "even")
            even_ref.set(qn("r:id"), rel)
            sect_pr.append(even_ref)

    @staticmethod
    def _parse_page_format(fmt: str) -> tuple[str, str, str]:
        """Parse format string → (prefix, PAGE instr_text, suffix).

        Placeholders: 1→Arabic, 一→Chinese, I→Roman upper, i→Roman lower.
        """
        for char, instr in [
            ("一", r" PAGE \* CHINESENUM "),
            ("I", r" PAGE \* ROMAN "),
            ("i", r" PAGE \* roman "),
            ("1", " PAGE "),
        ]:
            idx = fmt.find(char)
            if idx != -1:
                return fmt[:idx], instr, fmt[idx + len(char):]
        return fmt, "", ""

    def _write_page_number_para(self, para, align: str, pn, qn) -> None:
        from docx.oxml import OxmlElement

        align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
        para.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.CENTER)

        prefix, instr_text, suffix = self._parse_page_format(pn.format)
        font = pn.font_name
        size = Pt(pn.size_pt)

        if prefix:
            run = para.add_run(prefix)
            run.font.name = font
            run.font.size = size

        if instr_text:
            run = para.add_run()
            fld_begin = OxmlElement("w:fldChar")
            fld_begin.set(qn("w:fldCharType"), "begin")
            run._element.append(fld_begin)
            run.font.name = font
            run.font.size = size

            run2 = para.add_run()
            instr_el = OxmlElement("w:instrText")
            instr_el.set(qn("xml:space"), "preserve")
            instr_el.text = instr_text
            run2._element.append(instr_el)
            run2.font.name = font
            run2.font.size = size

            run3 = para.add_run()
            fld_end = OxmlElement("w:fldChar")
            fld_end.set(qn("w:fldCharType"), "end")
            run3._element.append(fld_end)
            run3.font.name = font
            run3.font.size = size

        if suffix:
            run = para.add_run(suffix)
            run.font.name = font
            run.font.size = size

    @staticmethod
    def _build_page_number_xml(p_elem, align: str, pn, qn, OxmlElement) -> None:
        """Build page number paragraph XML directly on an lxml element."""
        prefix, instr_text, suffix = WordWriter._parse_page_format(pn.format)
        font = pn.font_name
        sz_val = str(int(pn.size_pt * 2))  # half-points

        pPr = OxmlElement("w:pPr")
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), align)
        pPr.append(jc)
        p_elem.append(pPr)

        def _add_run(text=None, fld_char_type=None):
            r = OxmlElement("w:r")
            rPr = OxmlElement("w:rPr")
            rFonts = OxmlElement("w:rFonts")
            rFonts.set(qn("w:ascii"), font)
            rFonts.set(qn("w:hAnsi"), font)
            rFonts.set(qn("w:eastAsia"), font)
            rPr.append(rFonts)
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), sz_val)
            rPr.append(sz)
            szCs = OxmlElement("w:szCs")
            szCs.set(qn("w:val"), sz_val)
            rPr.append(szCs)
            r.append(rPr)
            if text is not None:
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = text
                r.append(t)
            if fld_char_type is not None:
                fld = OxmlElement("w:fldChar")
                fld.set(qn("w:fldCharType"), fld_char_type)
                r.append(fld)
            p_elem.append(r)

        if prefix:
            _add_run(text=prefix)
        if instr_text:
            _add_run(fld_char_type="begin")
            instr_r = OxmlElement("w:r")
            instr_el = OxmlElement("w:instrText")
            instr_el.set(qn("xml:space"), "preserve")
            instr_el.text = instr_text
            instr_r.append(instr_el)
            p_elem.append(instr_r)
            _add_run(fld_char_type="end")
        if suffix:
            _add_run(text=suffix)
